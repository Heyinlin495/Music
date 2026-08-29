"""
软件测试实习生简历模板生成脚本
基于音乐流媒体平台项目经历，生成专业Word格式简历
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os


def set_font(run, name='微软雅黑', size=None, bold=False, color=None):
    run.font.name = name
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'BBBBBB')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_title(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2B579A')
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run(title)
    set_font(run, size=13, bold=True, color=(0x1A, 0x1A, 0x2E))


def add_subtitle(doc, title, time=''):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, size=11, bold=True, color=(0x1A, 0x1A, 0x1A))
    if time:
        r2 = p.add_run(f'    {time}')
        set_font(r2, size=9, color=(0x88, 0x88, 0x88))


def add_info(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label)
    set_font(r, size=10, bold=True)
    r2 = p.add_run(value)
    set_font(r2, size=10)


def add_bullet(doc, text, bold_prefix=None, indent=Cm(0.6)):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = indent
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_font(r, size=10, bold=True)
        r2 = p.add_run(text)
        set_font(r2, size=10)
    else:
        r = p.add_run(text)
        set_font(r, size=10)


def create_resume():
    doc = Document()

    # ── Page setup ──
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # ── Default style ──
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.25

    # ══════════════════════════════════════
    # 姓名
    # ══════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run('[你的姓名]')
    set_font(run, size=22, bold=True, color=(0x1A, 0x1A, 0x2E))

    # 联系方式
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    contacts = [('电话：', '1XX-XXXX-XXXX'), ('邮箱：', 'xxxxxx@xxx.com'), ('微信：', 'xxxxxxxx')]
    for i, (label, value) in enumerate(contacts):
        if i > 0:
            sep = p.add_run('  |  ')
            set_font(sep, size=10, color=(0xAA, 0xAA, 0xAA))
        r = p.add_run(label)
        set_font(r, size=10, bold=True)
        r2 = p.add_run(value)
        set_font(r2, size=10)

    # 求职意向
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    intents = [('求职意向：', '软件测试实习生'), ('期望城市：', '[城市]'), ('到岗时间：', '随时 / [具体日期]')]
    for i, (label, value) in enumerate(intents):
        if i > 0:
            sep = p.add_run('  |  ')
            set_font(sep, size=10, color=(0xAA, 0xAA, 0xAA))
        r = p.add_run(label)
        set_font(r, size=10, bold=True, color=(0x2B, 0x57, 0x9A))
        r2 = p.add_run(value)
        set_font(r2, size=10, color=(0x2B, 0x57, 0x9A))

    add_divider(doc)

    # ══════════════════════════════════════
    # 教育背景
    # ══════════════════════════════════════
    add_section_title(doc, '教育背景')
    add_subtitle(doc, '[大学名称] — [专业名称]（本科）', '20XX.09 - 20XX.06')
    add_bullet(doc, 'GPA：X.XX / 4.0（或排名前 XX%）')
    add_bullet(doc, '相关课程：软件工程、软件测试、数据库原理、计算机网络、操作系统、数据结构与算法')
    add_bullet(doc, '[奖学金 / 荣誉，没有则删除此行]')

    add_divider(doc)

    # ══════════════════════════════════════
    # 专业技能
    # ══════════════════════════════════════
    add_section_title(doc, '专业技能')
    skills = [
        ('测试基础：', '掌握黑盒测试方法（等价类划分、边界值分析、判定表、场景法），能独立编写测试用例与缺陷报告'),
        ('测试工具：', '熟练使用 Jira / 禅道进行缺陷管理，了解 Fiddler / Postman 进行接口抓包与调试'),
        ('编程语言：', 'Python 基础扎实，能编写自动化测试脚本；了解 Java 基础语法'),
        ('数据库：', '熟练使用 MySQL，掌握单表 / 多表查询、子查询、聚合函数，能通过 SQL 验证数据一致性'),
        ('Linux：', '掌握常用命令（文件操作、日志查看、进程管理、权限配置），能在 Linux 环境下部署和排查问题'),
        ('版本控制：', '熟练使用 Git 进行分支管理、代码合并与冲突解决，了解 CI/CD 流程'),
        ('接口测试：', '了解 RESTful API 规范，能使用 Postman / Python requests 进行接口功能验证'),
    ]
    for label, value in skills:
        add_bullet(doc, value, bold_prefix=label)

    add_divider(doc)

    # ══════════════════════════════════════
    # 项目经历
    # ══════════════════════════════════════
    add_section_title(doc, '项目经历')

    # 项目 1
    add_subtitle(doc, '全栈音乐流媒体平台 — 测试负责人', '20XX.XX - 20XX.XX')

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run('项目简介：')
    set_font(r, size=10, bold=True)
    r2 = p.add_run('基于 Spring Boot + React 的自托管音乐流媒体平台，包含用户端和管理后台两个独立 SPA，支持音乐播放、歌单管理、用户管理等完整功能。')
    set_font(r2, size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('技术栈：')
    set_font(r, size=10, bold=True)
    r2 = p.add_run('Java / Spring Boot / React / MySQL / Docker / Nginx / JWT / Git')
    set_font(r2, size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run('我的职责与成果（STAR法则）：')
    set_font(r, size=10, bold=True)

    star = [
        ('S — 情境：', '项目开发完成后缺乏系统测试，核心功能（注册登录、歌曲上传、歌单管理）存在未覆盖的测试场景'),
        ('T — 任务：', '负责对后端 15+ 个 RESTful API 接口进行全面的功能测试与接口测试'),
        ('A — 行动：', '使用等价类划分和边界值分析设计测试用例，覆盖正常流程与异常场景（空参数、越权访问、重复提交等）；通过 Postman 逐个验证接口请求与响应；编写 SQL 查询语句校验数据库写入的准确性；使用 Git 管理测试代码，编写 20+ 个 JUnit 单元测试类，覆盖 Controller、Service、DTO、Security 等核心模块'),
        ('R — 结果：', '累计发现并提交缺陷 15+ 个，其中包含 3 个高优先级缺陷（未授权访问、数据校验缺失）；单元测试覆盖率达 75%+，有效拦截回归缺陷；项目通过 CI/CD 流水线自动执行测试，保障每次提交的质量'),
    ]
    for label, value in star:
        add_bullet(doc, value, bold_prefix=label)

    add_divider(doc)

    # ══════════════════════════════════════
    # 实习经历
    # ══════════════════════════════════════
    add_section_title(doc, '实习经历')
    add_subtitle(doc, '[公司名称] — 软件测试实习生', '20XX.XX - 20XX.XX')
    add_bullet(doc, '参与 [产品名称] 的功能测试，负责 [XX 模块] 的测试用例设计与执行，累计编写用例 XX+ 条')
    add_bullet(doc, '使用 [工具名称] 提交和跟踪缺陷，与开发团队协作推动缺陷修复，缺陷修复率达 XX%')
    add_bullet(doc, '协助编写测试报告，总结测试结果与风险项，为版本发布提供质量评估依据')

    add_divider(doc)

    # ══════════════════════════════════════
    # 校园经历
    # ══════════════════════════════════════
    add_section_title(doc, '校园经历')
    add_subtitle(doc, '[社团 / 组织名称] — [职位]', '20XX.XX - 20XX.XX')
    add_bullet(doc, '[描述你做了什么，突出组织能力、沟通能力或技术相关活动]')
    add_subtitle(doc, '[竞赛 / 活动名称]', '20XX.XX')
    add_bullet(doc, '[描述参与情况和成果，如蓝桥杯、数学建模、编程比赛等]')

    add_divider(doc)

    # ══════════════════════════════════════
    # 自我评价
    # ══════════════════════════════════════
    add_section_title(doc, '自我评价')
    evaluations = [
        '对软件测试有浓厚兴趣，系统学习过测试理论与方法，具备独立设计测试用例和执行测试的能力',
        '有完整的全栈项目测试实战经验，能从用户视角发现功能缺陷，也能通过接口和数据库进行深层验证',
        '具备良好的学习能力和自驱力，能快速上手新工具和新业务；注重细节，善于沟通协作',
        '认同质量是团队共同的责任，愿意在实际工作中持续提升测试技能和工程素养',
    ]
    for e in evaluations:
        add_bullet(doc, e)

    # ── Save ──
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '软件测试实习生简历模板.docx')
    doc.save(output_path)
    print(f'简历已生成: {output_path}')
    return output_path


if __name__ == '__main__':
    create_resume()
